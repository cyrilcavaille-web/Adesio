# -*- coding: utf-8 -*-
"""Export des lignes extraites au format d'import Excalibur ED36."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Colonnes A..S dans l'ordre exact de la maquette ED36.
# (obligatoire) = colonne requise par le parametrage d'import.
COLONNES = [
    ("N° cmde",          "num_cmde",        12, None,          True),
    ("Id",               "id_poste",         6, "@",           True),
    ("Produit",          "produit",         12, "@",           False),
    ("Qté",              "qte",             12, "#,##0.000",   False),
    ("Unité",            "unite",            8, None,          False),
    ("Coeff",            "coeff",           12, "#,##0.000",   False),
    ("Conditionnement",  "conditionnement", 16, None,          False),
    ("Qté/cond",         "qte_cond",        12, "#,##0.000",   False),
    ("PU HT",            "pu_ht",           12, "#,##0.000000", True),
    ("Remise",           "remise",           9, "#,##0.000",   False),
    ("Qté à recv.",      "qte_a_recevoir",  12, "#,##0.000",   True),
    ("Date confirmée",   "date_confirmee",  14, "DD/MM/YYYY",  True),
    ("Date AR",          "date_ar",         12, "DD/MM/YYYY",  False),
    ("N°AR fournisseur", "num_ar",          16, "@",           False),
    ("Réf. fournisseur", "ref_fournisseur", 16, "@",           False),
    ("Fabricant",        "fabricant",       26, None,          False),
    ("Réf. fabricant",   "ref_fabricant",   26, None,          False),
    ("Devise",           "devise",           8, None,          False),
    ("Commentaires",     "commentaires",    46, None,          False),
]

POLICE = "Arial"
BLEU = PatternFill("solid", fgColor="1F3864")
BLEU_CLAIR = PatternFill("solid", fgColor="D9E2F3")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
ROUGE = PatternFill("solid", fgColor="F8CBAD")
FIN = Side(style="thin", color="BFBFBF")
CADRE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)


def _niveau(ligne) -> str:
    txt = " ".join(ligne.anomalies)
    if "[BLOQUANT]" in txt:
        return "BLOQUANT"
    if "[A VERIFIER]" in txt or "[AVERTISSEMENT]" in txt:
        return "A CONTROLER"
    return "OK"


def exporter(lignes, chemin: Path) -> Path:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ---------------- Feuille 1 : le fichier d'import lui-meme ----------------
    ws = wb.active
    ws.title = "Import AR"
    for i, (titre, _, largeur, fmt, oblig) in enumerate(COLONNES, start=1):
        c = ws.cell(row=1, column=i, value=titre)
        c.font = Font(name=POLICE, bold=True, size=10, color="FFFFFF")
        c.fill = BLEU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CADRE
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    for r, ligne in enumerate(lignes, start=2):
        niveau = _niveau(ligne)
        for i, (_, attr, _, fmt, oblig) in enumerate(COLONNES, start=1):
            val = getattr(ligne, attr)
            c = ws.cell(row=r, column=i, value=val)
            c.font = Font(name=POLICE, size=10)
            c.border = CADRE
            if fmt:
                c.number_format = fmt
            if val in (None, "") and oblig:
                c.fill = ROUGE
            elif niveau == "A CONTROLER" and oblig:
                c.fill = ORANGE
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{max(2, len(lignes) + 1)}"

    # ---------------- Feuille 2 : journal de controle ----------------
    wc = wb.create_sheet("Contrôles")
    entetes = ["Fichier", "Fournisseur", "Méthode", "N° cmde", "Poste",
               "Statut", "Confiance", "Anomalies"]
    for i, t in enumerate(entetes, start=1):
        c = wc.cell(row=1, column=i, value=t)
        c.font = Font(name=POLICE, bold=True, size=10, color="FFFFFF")
        c.fill = BLEU
        c.border = CADRE
    for i, w in enumerate([34, 24, 9, 14, 7, 14, 10, 78], start=1):
        wc.column_dimensions[get_column_letter(i)].width = w
    for r, l in enumerate(lignes, start=2):
        niveau = _niveau(l)
        valeurs = [l.fichier, l.fournisseur, l.methode, l.num_cmde, l.id_poste,
                   niveau, l.confiance, " | ".join(l.anomalies) or "-"]
        for i, v in enumerate(valeurs, start=1):
            c = wc.cell(row=r, column=i, value=v)
            c.font = Font(name=POLICE, size=10)
            c.border = CADRE
            c.alignment = Alignment(vertical="top", wrap_text=(i == 8))
            if i == 6:
                c.fill = {"BLOQUANT": ROUGE, "A CONTROLER": ORANGE}.get(niveau, BLEU_CLAIR)
            if i == 7:
                c.number_format = "0.00"
    wc.freeze_panes = "A2"

    # ---------------- Feuille 3 : parametrage ED36 applique ----------------
    wp = wb.create_sheet("Paramétrage ED36")
    wp.column_dimensions["A"].width = 42
    wp.column_dimensions["B"].width = 20
    wp.column_dimensions["C"].width = 62
    params = [
        ("Paramètre d'import", "Valeur", "Commentaire"),
        ("Nombre de lignes d'entête à ignorer", 1, "ligne 1 = titres de colonnes"),
        ("Écart toléré date confirmée / souhaitée (j)", 0,
         "au-delà, le traitement lève un avertissement"),
        ("Date confirmée à utiliser en cas d'allocation", "00/00/0000", "paramètre Excalibur"),
        ("", "", ""),
        ("Mapping colonne → zone Excalibur", "Colonne", "Obligatoire"),
    ]
    for titre, val, com in params:
        wp.append([titre, val, com])
    for i, (titre, _, _, _, oblig) in enumerate(COLONNES, start=1):
        wp.append([titre, get_column_letter(i), "OUI" if oblig else "facultative"])
    for row in wp.iter_rows():
        for c in row:
            c.font = Font(name=POLICE, size=10,
                          bold=c.row in (1, 6) or (c.column == 3 and c.value == "OUI"))
            c.border = CADRE
            if c.row in (1, 6):
                c.fill = BLEU_CLAIR

    wb.save(chemin)
    return chemin
